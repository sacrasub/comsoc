# -*- coding: utf-8 -*-
"""
Script para geração do Manual Operacional e Guia Passo a Passo do Sistema ComSoc CFT
Capitania Fluvial de Tabatinga • Marinha do Brasil
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Define a cor de fundo de uma célula."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Define margens internas da célula."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, items, border_color="002B5B", bg_color="F1F5F9"):
    """Cria uma caixa de destaque visual com borda esquerda colorida."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    first = True
    for item in items:
        # Se for uma lista de runs (tuplas) para a mesma linha/parágrafo
        if isinstance(item, list):
            if not first:
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
            first = False
            for rspec in item:
                text = str(rspec[0])
                bold = rspec[1] if len(rspec) > 1 else False
                italic = rspec[2] if len(rspec) > 2 else False
                color = rspec[3] if len(rspec) > 3 else None
                run = p.add_run(text)
                run.font.name = 'Calibri'
                run.font.size = Pt(10)
                run.font.bold = bold
                run.font.italic = italic
                if color:
                    run.font.color.rgb = color
                else:
                    run.font.color.rgb = RGBColor(30, 41, 59)
        elif isinstance(item, tuple):
            text = str(item[0])
            bold = item[1] if len(item) > 1 else False
            italic = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 else None
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.bold = bold
            run.font.italic = italic
            if color:
                run.font.color.rgb = color
            else:
                run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            if not first:
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
            first = False
            run = p.add_run(str(item))
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(30, 41, 59)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table_header(row, col_widths, bg_color="002B5B"):
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

def format_table_row(row, col_widths, bg_color="FFFFFF", align_left=True):
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            if not align_left and i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(30, 41, 59)

def build_manual_docx():
    doc = Document()
    
    # Configuração de Margens (2.0 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Cabeçalho da página
        hdr = section.header
        hp = hdr.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("MARINHA DO BRASIL • CAPITANIA FLUVIAL DE TABATINGA • SEÇÃO DE COMSOC")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Rodapé da página
        ftr = section.footer
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Manual Operacional e Diretrizes ComSoc CFT • Último Trimestre de 2026")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # ═════════════════════════════════════════════════════════════════════════
    # CAPA / CABEÇALHO MILITAR NAVAL
    # ═════════════════════════════════════════════════════════════════════════
    p_topo = doc.add_paragraph()
    p_topo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topo.paragraph_format.space_before = Pt(10)
    p_topo.paragraph_format.space_after = Pt(2)
    r1 = p_topo.add_run("MARINHA DO BRASIL\n")
    r1.font.name = 'Calibri'
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 43, 91)
    
    r2 = p_topo.add_run("COMANDO DO 9º DISTRITO NAVAL\n")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(11)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(26, 77, 128)
    
    r3 = p_topo.add_run("CAPITANIA FLUVIAL DE TABATINGA\n")
    r3.font.name = 'Calibri'
    r3.font.size = Pt(12)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0, 43, 91)
    
    r4 = p_topo.add_run("SEÇÃO DE COMUNICAÇÃO SOCIAL (ComSoc) • DIVISÃO 40 (EMBARCAÇÕES E VIATURAS)")
    r4.font.name = 'Calibri'
    r4.font.size = Pt(9.5)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(100, 116, 139)

    # Linha divisória ornamental
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(16)
    r_div = p_div.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    r_div.font.name = 'Calibri'
    r_div.font.size = Pt(8)
    r_div.font.color.rgb = RGBColor(212, 175, 55)

    # Título Principal do Manual
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_after = Pt(4)
    r_tit = p_tit.add_run("MANUAL OPERACIONAL E DIRETRIZES DO SISTEMA COMSOC CFT")
    r_tit.font.name = 'Calibri'
    r_tit.font.size = Pt(18)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(0, 43, 91)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("Guia Prático Passo a Passo para Capacitação, Delegação Funcional e Tomada de Decisões\n(Período Crítico de Outubro a Dezembro de 2026)")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Caixa de Metadados Oficiais
    add_callout(doc, [
        [("ENCARREGADO COMSOC / AUTOR: ", True, False, RGBColor(0, 43, 91)), ("SO-CI-SB Sacramento (Suboficial-Mor da CFT)", False, False, None)],
        [("AUXILIARES DESIGNADOS: ", True, False, RGBColor(0, 43, 91)), ("SG-MO Matheus (Operacional/Div-40) e CB-RM2-PD Marla (Administrativa/OCB)", False, False, None)],
        [("JURISDIÇÃO: ", True, False, RGBColor(0, 43, 91)), ("Tríplice Fronteira Amazônica • Tabatinga e Benjamin Constant - AM", False, False, None)],
        [("FINALIDADE: ", True, False, RGBColor(0, 43, 91)), ("Orientar o Encarregado e capacitar a equipe para operação autônoma durante ausências, garantindo continuidade e excelência.", False, True, None)]
    ], border_color="002B5B", bg_color="F8FAFC")

    # ═════════════════════════════════════════════════════════════════════════
    # SEÇÃO 1: APRESENTAÇÃO E PROPÓSITO
    # ═════════════════════════════════════════════════════════════════════════
    h1 = doc.add_heading("1. APRESENTAÇÃO E OBJETIVO DO SISTEMA", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0, 43, 91)

    doc.add_paragraph(
        "O Sistema ComSoc CFT é uma plataforma integrada desenvolvida especificamente para as necessidades estratégicas e operacionais "
        "da Capitania Fluvial de Tabatinga. O sistema centraliza o controle do Diretório de Autoridades civis e militares do Alto Solimões, "
        "a Agenda Oficial de Eventos com acompanhamento de presenças (RSVP), o checklist normativo de cerimonial (EMA-136 e EMA-860), "
        "e o novo módulo de Gestão de Processos, Delegação e Transição."
    )
    doc.add_paragraph(
        "Este manual tem dois objetivos primordiais:\n"
        "1. Servir como roteiro de consulta e comando para o Suboficial-Mor (SO Sacramento), assegurando o controle das decisões e a fiscalização dos prazos;\n"
        "2. Servir como instrumento pedagógico para capacitar o SG-MO Matheus e a CB-RM2-PD Marla, capacitando-os a conduzir as rotinas da seção com total autonomia e disciplina militar durante os períodos de ausência do titular e pavimentando a futura transição da seção para a reserva."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SEÇÃO 2: ESTRUTURA DE DELEGAÇÃO FUNCIONAL
    # ═════════════════════════════════════════════════════════════════════════
    h2 = doc.add_heading("2. ESTRUTURA DE DELEGAÇÃO FUNCIONAL E TRANSIÇÃO", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0, 43, 91)

    doc.add_paragraph(
        "A compatibilização da rotina da Seção de Comunicação Social (ComSoc) com os encargos da Divisão 40 (Embarcações e Viaturas) "
        "estabelece uma separação funcional clara, objetiva e complementar entre os militares da equipe:"
    )

    # Tabela da Equipe
    tbl_equipe = doc.add_table(rows=1, cols=3)
    tbl_equipe.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_equipe.autofit = False
    col_w = [Inches(1.8), Inches(1.8), Inches(2.9)]
    
    hdr_cells = tbl_equipe.rows[0].cells
    hdr_cells[0].paragraphs[0].add_run("Militar / Graduação")
    hdr_cells[1].paragraphs[0].add_run("Função no Sistema")
    hdr_cells[2].paragraphs[0].add_run("Atribuições Principais & Foco")
    format_table_header(tbl_equipe.rows[0], col_w)

    dados_equipe = [
        ("SO-CI-SB SACRAMENTO", "Encarregado ComSoc\nSuboficial-Mor (SOMor)", "• Supervisão Geral da rotina da ComSoc e cerimonial naval\n• Despacho permanente com o Comandante e Imediato\n• Planejamento estratégico e mentoria contínua da equipe\n• Ponto de controle e aprovação final de matérias/releases"),
        ("SG-MO MATHEUS", "Auxiliar Operacional\n(ComSoc & Div-40)", "• Foco Logístico e Meios: conciliar manutenção das lanchas (Div-40) com apoio a eventos e travessias náuticas (ex: Benjamin Constant)\n• Apoio Técnico de Som e Palanque: montagem e teste de sonorização, microfones e projeção\n• Acervo Audiovisual: registro fotográfico de alta resolução e salvamento padronizado no servidor da Secom"),
        ("CB-RM2-PD MARLA", "Auxiliar Administrativa\ne de Conteúdo", "• Operação Cisne Branco (OCB): triagem de redações digitais, pastas físicas por escola e contato com diretorias\n• Redação Institucional: matérias da Bússola Amazônica, minutas de ofícios e roteiros de locução (vogal)\n• Interlocução Externa: canal direto com SEDUC e secretarias municipais de Tabatinga e Benjamin Constant")
    ]

    for i, d in enumerate(dados_equipe):
        row = tbl_equipe.add_row()
        row.cells[0].paragraphs[0].add_run(d[0])
        row.cells[1].paragraphs[0].add_run(d[1])
        row.cells[2].paragraphs[0].add_run(d[2])
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        format_table_row(row, col_w, bg_color=bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ═════════════════════════════════════════════════════════════════════════
    # SEÇÃO 3: GUIA PASSO A PASSO DE OPERAÇÃO DO SISTEMA
    # ═════════════════════════════════════════════════════════════════════════
    h3 = doc.add_heading("3. GUIA PASSO A PASSO: UTILIZAÇÃO DE CADA MÓDULO", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0, 43, 91)

    # Subseção 3.1
    h31 = doc.add_heading("3.1. Como Acessar o Sistema", level=2)
    h31.paragraph_format.space_before = Pt(8)
    h31.paragraph_format.space_after = Pt(4)
    for r in h31.runs:
        r.font.color.rgb = RGBColor(26, 77, 128)

    doc.add_paragraph(
        "1. No computador da Seção ou em qualquer máquina autorizada na rede, abra o navegador Google Chrome ou Microsoft Edge;\n"
        "2. Acesse o endereço do painel (ou abra o arquivo index.html no servidor compartilhado da CFT);\n"
        "3. Faça o login utilizando sua conta Google institucional ou e-mail cadastrado na whitelist oficial;\n"
        "4. O sistema reconhecerá automaticamente seu perfil (Super Admin, Administrador ou Operador) e carregará as permissões correspondentes."
    )

    # Subseção 3.2
    h32 = doc.add_heading("3.2. Módulo 'Diretório de Autoridades' & Controle de Presença (RSVP)", level=2)
    h32.paragraph_format.space_before = Pt(8)
    h32.paragraph_format.space_after = Pt(4)
    for r in h32.runs:
        r.font.color.rgb = RGBColor(26, 77, 128)

    doc.add_paragraph(
        "Este módulo é o coração institucional da Capitania para eventos solenes:\n"
        "• Barra de Evento Ativo: No topo da tela, você sempre verá qual evento está selecionado (ex: 'Aniversário de Ativação da CFT'). Todas as marcações de RSVP feitas nos contatos pertencem a este evento;\n"
        "• Filtro por Categoria: Clique nos cards superiores para filtrar autoridades por categoria (Segurança/Militar, Executivo/Prefeituras, Judiciário ou SOAMAR);\n"
        "• Marcação Individual de Presença: Na lista de contatos, clique no botão de status da autoridade para alternar entre Pendente (cinza), Convite Enviado (azul), Confirmado (verde) ou Recusado (vermelho);\n"
        "• Ações em Lote (Disparos / Confirmações em Massa): Marque a caixa de seleção de várias autoridades. A barra flutuante preta surgirá na parte inferior da tela, permitindo alterar o status de presença de todos de uma só vez;\n"
        "• Exportação de Relatórios: Utilize os botões do diretório para gerar relatórios em PDF ou planilhas Excel (XLSX) formatadas para o Gabinete do Comandante."
    )

    # Subseção 3.3
    h33 = doc.add_heading("3.3. Módulo 'Agenda de Eventos'", level=2)
    h33.paragraph_format.space_before = Pt(8)
    h33.paragraph_format.space_after = Pt(4)
    for r in h33.runs:
        r.font.color.rgb = RGBColor(26, 77, 128)

    doc.add_paragraph(
        "• Timeline Cronológica: Exibe todos os eventos cadastrados em ordem de data, informando local, horário e uniforme regulamentar;\n"
        "• Botão 'Tornar Evento Ativo': Ao clicar neste botão em qualquer evento da agenda, ele passa a ser o evento ativo do sistema, atualizando o banner e o RSVP do Diretório;\n"
        "• Criação e Edição: Utilize o formulário lateral para criar novos eventos ou atualizar dados (data, horário, uniforme e descrição)."
    )

    # Subseção 3.4
    h34 = doc.add_heading("3.4. Módulo 'Gerenciamento ComSoc' (Checklist Normativo EMA-136/860)", level=2)
    h34.paragraph_format.space_before = Pt(8)
    h34.paragraph_format.space_after = Pt(4)
    for r in h34.runs:
        r.font.color.rgb = RGBColor(26, 77, 128)

    doc.add_paragraph(
        "O módulo Gerenc. ComSoc organiza o cerimonial em 3 fases regulamentares da Marinha:\n"
        "1. Fase 1: Planejamento (Lista de convidados, Roteiro do Vogal, Marcação de dispositivo, Teste de som/projetor, Aviso de pauta, Briefing da equipe);\n"
        "2. Fase 2: Execução (Recepção de imprensa, Mediação de entrevistas, Cobertura fotográfica estratégica sem obstruir autoridades);\n"
        "3. Fase 3: Pós-Evento (Elaboração de release institucional, Envio ao Com9ºDN, Nota no Plano do Dia, Clipping de notícias e Avaliação de resultados);\n"
        "• Ações Rápidas: Ao lado de cada tarefa, há botões que abrem diretamente os geradores de Roteiro do Vogal, cadastro de Clipping de mídia e registro de matérias."
    )

    # Subseção 3.5 - O NOVO MÓDULO
    h35 = doc.add_heading("3.5. Módulo 'Gestão de Processos & Delegação' (Aba 15 DIAS)", level=2)
    h35.paragraph_format.space_before = Pt(8)
    h35.paragraph_format.space_after = Pt(4)
    for r in h35.runs:
        r.font.color.rgb = RGBColor(26, 77, 128)

    doc.add_paragraph(
        "A nova aba de Gestão de Processos reúne as ferramentas executivas para acompanhamento em tempo real das metas imediatas "
        "e tomada de decisão com o Comando. Possui 4 visões interativas selecionáveis:"
    )

    add_callout(doc, [
        [("VISÃO 1: MATRIZ DOS PRÓXIMOS 15 DIAS (01 A 16 DE OUTUBRO)", True, False, RGBColor(0, 43, 91))],
        [("• O que é: ", True, False, None), ("Painel com as 6 metas prioritárias durante o período de viagem/ausência do SOMor.", False, False, None)],
        [("• Como operar: ", True, False, None), ("Ao concluir ou avançar uma tarefa, clique no botão correspondente (Pendente, Andamento ou Concluído). O sistema atualiza a barra de progresso imediatamente.", False, False, None)],
        [("• Anotações de Despacho: ", True, False, None), ("Clique no ícone de lápis para registrar números de protocolo, nomes de diretores escolares contatados ou observações de despacho.", False, False, None)],
        [("• Checklist de Passagem de Serviço: ", True, False, None), ("Na parte inferior, marque os 3 passos práticos (Briefing Inicial, Despacho Imediato e Ponto de Controle na Volta).", False, False, None)]
    ], border_color="F59E0B", bg_color="FFFBEB")

    add_callout(doc, [
        [("VISÃO 2: CALENDÁRIO & DECISÕES DO COMANDO (OUTUBRO A DEZEMBRO)", True, False, RGBColor(0, 43, 91))],
        [("• O que é: ", True, False, None), ("Exibição dos 6 eventos oficiais do último trimestre com seus respectivos pontos de tomada de decisão.", False, False, None)],
        [("• Como operar: ", True, False, None), ("Nos eventos que possuem alternativas (ex: Ativação da CFT: Opção A bordo vs Opção CAMATA), selecione a opção aprovada pelo Comandante. O sistema salva a escolha com data e nome de quem registrou.", False, False, None)],
        [("• Botão 'Gerenciar RSVP no Diretório': ", True, False, None), ("Clique para saltar diretamente ao Diretório com aquele evento selecionado e marcar as presenças das autoridades convidadas.", False, False, None)]
    ], border_color="3B82F6", bg_color="EFF6FF")

    add_callout(doc, [
        [("VISÃO 3: ESTRUTURA DE DELEGAÇÃO & EQUIPE", True, False, RGBColor(0, 43, 91))],
        [("• O que é: ", True, False, None), ("Organograma vivo demonstrando a cadeia funcional (SO Sacramento no topo, SG Matheus e CB Marla como auxiliares).", False, False, None)],
        [("• Contador de Metas: ", True, False, None), ("Mostra dinamicamente quantas tarefas cada militar possui e quantas já foram entregues.", False, False, None)]
    ], border_color="10B981", bg_color="ECFDF5")

    add_callout(doc, [
        [("VISÃO 4: QUADRO KANBAN OPERACIONAL", True, False, RGBColor(0, 43, 91))],
        [("• O que é: ", True, False, None), ("Visualização em colunas ('A Fazer / Pendente', 'Em Andamento', 'Concluído').", False, False, None)],
        [("• Como operar: ", True, False, None), ("Utilize as setas e botões nos cartões para mover tarefas entre colunas conforme o expediente avança.", False, False, None)]
    ], border_color="6366F1", bg_color="EEF2FF")

    # ═════════════════════════════════════════════════════════════════════════
    # SEÇÃO 4: MATRIZ CRÍTICA DOS 15 DIAS (01 A 16 DE OUTUBRO DE 2026)
    # ═════════════════════════════════════════════════════════════════════════
    h4 = doc.add_heading("4. MATRIZ DE GESTÃO: PRAZOS IMEDIATOS (01 A 16/OUT)", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0, 43, 91)

    doc.add_paragraph(
        "Esta matriz detalha exatamente as obrigações e prazos fatais da equipe durante a ausência temporária do Suboficial-Mor. "
        "O cumprimento estrito destes prazos é essencial para não comprometer a remessa das redações da OCB ao 9º Distrito Naval "
        "nem a organização do aniversário de ativação da Capitania:"
    )

    tbl_15 = doc.add_table(rows=1, cols=5)
    tbl_15.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_15.autofit = False
    col_15 = [Inches(0.5), Inches(2.3), Inches(1.2), Inches(1.3), Inches(1.2)]
    
    h15 = tbl_15.rows[0].cells
    h15[0].paragraphs[0].add_run("#")
    h15[1].paragraphs[0].add_run("Meta / Tarefa")
    h15[2].paragraphs[0].add_run("Prazo Crítico")
    h15[3].paragraphs[0].add_run("Responsável")
    h15[4].paragraphs[0].add_run("Status Inicial")
    format_table_header(tbl_15.rows[0], col_15)

    metas_data = [
        ("1", "Encerramento da Coleta da OCB\n(Centralizar redações digitais e contato escolas)", "Até 09/10/2026", "CB Marla", "PENDENTE"),
        ("2", "Triagem das Redações Digitais\n(Conferir fichas, temas e notas da banca)", "05 a 12/10/2026", "CB Marla", "PENDENTE"),
        ("3", "Montagem das Pastas Físicas da OCB\n(Separar por colégios de TBT e BC)", "13 a 16/10/2026", "SG Matheus /\nCB Marla", "PENDENTE"),
        ("4", "Minuta de Ofícios Institucionais\n(Eventos de 26/10 e 06/11)", "Até 14/10/2026", "CB Marla", "PENDENTE"),
        ("5", "Levantamento Técnico e Teste de Som\n(Sonorização, microfones, telão auditório/Praça)", "Até 15/10/2026", "SG Matheus\n(Div-40)", "PENDENTE"),
        ("6", "Lembrete Relatório Com9ºDN\n(Deixar pronto para envio final)", "Pronto em 16/10\n(Envio 31/10)", "CB Marla", "PENDENTE")
    ]

    for i, m in enumerate(metas_data):
        row = tbl_15.add_row()
        row.cells[0].paragraphs[0].add_run(m[0])
        row.cells[1].paragraphs[0].add_run(m[1])
        row.cells[2].paragraphs[0].add_run(m[2])
        row.cells[3].paragraphs[0].add_run(m[3])
        row.cells[4].paragraphs[0].add_run(m[4])
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        format_table_row(row, col_15, bg_color=bg, align_left=False)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Passos Práticos de Passagem de Serviço
    doc.add_heading("Roteiro Prático de Passagem de Serviço antes da Partida do SOMor:", level=2)
    doc.add_paragraph(
        "1. Briefing Inicial com a Equipe (01/10): Reunir o SG Matheus e a CB Marla na Seção. Repassar as senhas de acesso aos canais digitais da ComSoc, a pasta compartilhada no servidor local e a agenda telefônica com os gestores escolares das redes SEDUC e municipal;\n"
        "2. Despacho Formal com o Imediato (02/10): Apresentar a escala de substituição temporária na ComSoc durante os 15 dias, assegurando o respaldo institucional da CB Marla para que suas atividades na Seção de Comunicação não sofram interferência de outras rotinas administrativas;\n"
        "3. Ponto de Controle no Retorno (19/10 - Segunda-feira pós-viagem): Reunião matutina de avaliação no primeiro dia de retorno para revisar o lote de redações selecionadas para envio ao 9º Distrito Naval e aprovar os roteiros finais da cerimônia de ativação de 26 de outubro."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SEÇÃO 5: CALENDÁRIO OFICIAL & DIRETRIZES DE DECISÃO
    # ═════════════════════════════════════════════════════════════════════════
    h5 = doc.add_heading("5. CALENDÁRIO OFICIAL & DIRETRIZES DE DECISÃO (OUT A DEZ)", level=1)
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)
    for r in h5.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0, 43, 91)

    doc.add_paragraph(
        "Diretrizes e justificativas regulamentares de cada um dos 6 eventos oficiais sincronizados no sistema:"
    )

    eventos_detalhes = [
        {
            "num": "5.1",
            "nome": "26/10 (Seg) — Aniversário de Ativação da Capitania Fluvial de Tabatinga",
            "horario": "08h00: Pátio Principal (Cerimônia Militar) | 09h00: Praça D’Armas (Café)",
            "uniforme": "5.5 (Tropa e Oficiais) / 6.4 (Conforme Ordem de Serviço)",
            "justificativa": "Cerimônia Militar a bordo com leitura da Ordem do Dia alusiva ao histórico da Capitania e condecorações internas.",
            "decisao": "PONTO DE DECISÃO DO COMANDO:\n"
                       "• Opção A (Recomendada pela ComSoc): Café da manhã especial / buffet na Praça D’Armas logo após a formatura do dia 26/10. Não compromete o expediente, a segurança orgânica nem o atendimento do GAP.\n"
                       "• Opção B (CAMATA Família Naval): Confraternização com churrasco e patrocínio das VCB no CAMATA no sábado que antecede (24/10), evitando esvaziar a rotina na segunda-feira útil.",
            "acoes": "• Montagem e teste de som na Praça D\'Armas (SG Matheus)\n• Cobertura fotográfica em alta resolução dos condecorados (SG Matheus)\n• Expedição de convites e matéria para a Bússola Amazônica (CB Marla)"
        },
        {
            "num": "5.2",
            "nome": "06/11 (Sex) — Dia Nacional do Amigo da Marinha",
            "horario": "10h00: Sede da CFT / Auditório Climatizado",
            "uniforme": "5.5 (Militar) / Passeio Completo (Civis)",
            "justificativa": "Cerimônia solene de imposição da Medalha 'Amigo da Marinha' a autoridades e personalidades locais do Alto Solimões e membros da SOAMAR.",
            "decisao": "PONTO DE DECISÃO DO COMANDO: Aprovação prévia da lista de agraciados e definição do formato de recepção no salão nobre.",
            "acoes": "• Expedição dos convites institucionais com 15 dias de antecedência (impreterivelmente até 22/10) — CB Marla\n• Contato prévio com a presidência da SOAMAR Alto Solimões para confirmação de presença — CB Marla\n• Roteiro de entrega de medalhas e diplomas com a Seção de Pessoal — CB Marla / SO Sacramento\n• Sonorização e projeção de vídeo institucional no auditório — SG Matheus"
        },
        {
            "num": "5.3",
            "nome": "13/11 (Sex) — Cerimônia Cívica na E. E. Almirante Tamandaré",
            "horario": "08h30: Quadra Coberta da Escola",
            "uniforme": "5.5 (Guarda de Honra e Representação)",
            "justificativa": "A sexta-feira (13/11) fecha com chave de ouro a semana pedagógica da escola antes do feriado da República (15/11), mobilizando todos os turnos escolares.",
            "decisao": "DIRETRIZES: Dispositivo de Guarda de Honra e representação da OM.",
            "acoes": "• Alinhamento com a diretoria pedagógica na última semana de outubro — CB Marla\n• Roteiro alusivo ao Patrono da Marinha e doação de material institucional/Bandeira — CB Marla\n• Som portátil e registro fotográfico dos estudantes com a tropa — SG Matheus"
        },
        {
            "num": "5.4",
            "nome": "19/11 (Qui) — Dia da Bandeira",
            "horario": "12h00 RIGOROSO (Meio-dia): Pátio Principal da CFT",
            "uniforme": "5.5",
            "justificativa": "Cumprimento estrito do Cerimonial Naval pontualmente ao meio-dia, com içamento extraordinário da Bandeira Nacional e incineração solene das bandeiras inservíveis.",
            "decisao": "DIRETRIZES: Presença de representações de escolas cívico-militares de Tabatinga.",
            "acoes": "• Convite formal a alunos de escolas locais para assistirem ao cerimonial cívico — CB Marla\n• Separação e preparo com o Contramestre das bandeiras inservíveis para incineração — SG Matheus\n• Registro fotográfico e cobertura da salva de tiros — SG Matheus"
        },
        {
            "num": "5.5",
            "nome": "11/12 (Sex) — Dia do Marinheiro (Data Magna 13/DEZ)",
            "horario": "10h00: Pátio de Formatura da CFT",
            "uniforme": "5.5 / 3.3 (Conforme diretriz do Comando)",
            "justificativa": "O dia 13 de dezembro de 2026 cai em um domingo. O cerimonial militar oficial com presença de autoridades e da sociedade é antecipado para a última sexta-feira útil (11/12).",
            "decisao": "PONTO DE DECISÃO DO COMANDO: Definição do uniforme de gala (3.3) para agraciados da Medalha Mérito Marinheiro e Oficiais da mesa, ou 5.5 para toda a tropa devido ao clima amazônico.",
            "acoes": "• Formatura solene com imposição da Medalha Mérito Marinheiro e promoção de Praças\n• Leitura da Ordem do Dia do Comandante da Marinha — Vogal / CB Marla\n• Registro completo no Livro do Estabelecimento da CFT — SG Matheus"
        },
        {
            "num": "5.6",
            "nome": "18/12 (Sex) — Cerimônia de Premiação da OCB 2026",
            "horario": "10h00: Auditório da CFT",
            "uniforme": "5.5 / Passeio",
            "justificativa": "Fechamento solene do ciclo letivo antes do recesso escolar de fim de ano, coroando os alunos vencedores de Tabatinga e Benjamin Constant.",
            "decisao": "PONTO CRÍTICO DE DECISÃO & LOGÍSTICA (Div-40):\n"
                       "• Escala de lancha da Capitania da Div-40 (SG Matheus) para buscar em Benjamin Constant os alunos e professores premiados (escolas CETI, Graziela e Imaculada Conceição) e retorná-los em segurança.",
            "acoes": "• Coordenação náutica e controle de combustível para a travessia fluvial — SG Matheus (Div-40)\n• Organização dos certificados, medalhas e kits de premiação — CB Marla\n• Recepção no portaló e condução ao auditório — Equipe ComSoc\n• Produção de matéria e informe imediato ao Com9ºDN — CB Marla"
        }
    ]

    for ev in eventos_detalhes:
        h_ev = doc.add_heading(f"{ev['num']}. {ev['nome']}", level=2)
        h_ev.paragraph_format.space_before = Pt(10)
        h_ev.paragraph_format.space_after = Pt(2)
        for r in h_ev.runs:
            r.font.color.rgb = RGBColor(0, 43, 91)
            
        p_info = doc.add_paragraph()
        p_info.paragraph_format.space_after = Pt(2)
        r_h = p_info.add_run(f"Horário & Local: ")
        r_h.bold = True
        p_info.add_run(f"{ev['horario']}\n")
        r_u = p_info.add_run(f"Uniforme Regulamentar: ")
        r_u.bold = True
        p_info.add_run(f"{ev['uniforme']}\n")
        r_j = p_info.add_run(f"Justificativa Operacional: ")
        r_j.bold = True
        p_info.add_run(f"{ev['justificativa']}")

        add_callout(doc, [
            [("DECISÃO / DIRETRIZ ESTRATÉGICA: \n", True, False, RGBColor(0, 43, 91)), (ev['decisao'], False, False, None)],
            [("\nAÇÕES COMSOC NECESSÁRIAS: \n", True, False, RGBColor(26, 77, 128)), (ev['acoes'], False, False, None)]
        ], border_color="002B5B", bg_color="F8FAFC")

    # ═════════════════════════════════════════════════════════════════════════
    # SEÇÃO 6: OPERAÇÃO OFFLINE, SINCRONIZAÇÃO E SEGURANÇA
    # ═════════════════════════════════════════════════════════════════════════
    h6 = doc.add_heading("6. OPERAÇÃO OFFLINE, SINCRONIZAÇÃO E SEGURANÇA", level=1)
    h6.paragraph_format.space_before = Pt(16)
    h6.paragraph_format.space_after = Pt(6)
    for r in h6.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0, 43, 91)

    doc.add_paragraph(
        "Dada a realidade de conectividade da Tríplice Fronteira Amazônica, o sistema foi projetado com alta tolerância a falhas:\n"
        "1. Funcionamento Offline Total (LocalStorage): Caso a internet caia, todas as alterações de RSVP, metas e decisões do Comando continuarão sendo salvas localmente no navegador;\n"
        "2. Sincronização Automática em Nuvem: Assim que a conexão for reestabelecida, o banco de dados Firebase Firestore recebe as atualizações sem perda de informações;\n"
        "3. Botão 'Sincronizar Agenda Oficial': Caso algum operador esteja com a lista desatualizada, basta clicar neste botão no cabeçalho para recarregar os 6 eventos normativos;\n"
        "4. Backup em Arquivo JSON: Recomenda-se que a CB Marla realize, toda sexta-feira, o download do backup local (menu 'Configurações' -> 'Backup Local JSON'), salvando na pasta da ComSoc no servidor da Capitania."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SEÇÃO 7: TERMO DE CIÊNCIA E COMPROMISSO OPERACIONAL
    # ═════════════════════════════════════════════════════════════════════════
    h7 = doc.add_heading("7. TERMO DE CIÊNCIA E COMPROMISSO OPERACIONAL", level=1)
    h7.paragraph_format.space_before = Pt(16)
    h7.paragraph_format.space_after = Pt(6)
    for r in h7.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0, 43, 91)

    doc.add_paragraph(
        "Declaramos ter tomado pleno conhecimento das diretrizes, prazos e responsabilidades funcionais contidas neste manual, "
        "comprometendo-nos a manter a operação da Seção de Comunicação Social e da Divisão 40 no mais alto grau de presteza e disciplina militar."
    )

    p_data = doc.add_paragraph()
    p_data.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_data.paragraph_format.space_before = Pt(15)
    p_data.paragraph_format.space_after = Pt(25)
    r_dt = p_data.add_run("Tabatinga - AM, 01 de outubro de 2026.")
    r_dt.font.bold = True

    # Tabela de Assinaturas
    tbl_ass = doc.add_table(rows=2, cols=2)
    tbl_ass.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ass.autofit = False
    col_ass = [Inches(3.2), Inches(3.2)]

    row0 = tbl_ass.rows[0]
    row0.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row0.cells[0].paragraphs[0].add_run("_________________________________________\nSO-CI-SB CRISTIANO SACRAMENTO\nEncarregado da ComSoc • Suboficial-Mor")
    
    row0.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row0.cells[1].paragraphs[0].add_run("_________________________________________\nSG-MO MATHEUS\nAuxiliar Operacional (ComSoc / Div-40)")

    row1 = tbl_ass.rows[1]
    row1.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row1.cells[0].paragraphs[0].add_run("\n\n_________________________________________\nCB-RM2-PD MARLA\nAuxiliar Administrativa e de Conteúdo")

    row1.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row1.cells[1].paragraphs[0].add_run("\n\nVISTO DO COMANDO:\n\n_________________________________________\nCOMANDANTE / IMEDIATO\nCapitania Fluvial de Tabatinga")

    for r in tbl_ass.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.5)
                    run.font.bold = True

    # Salvar documento
    out_path = os.path.abspath(r"c:\Projetos\Marinha\ComSoc\Manual_Operacional_ComSoc_CFT_2026.docx")
    doc.save(out_path)
    print(f"Sucesso! Documento salvo em: {out_path}")

if __name__ == "__main__":
    build_manual_docx()
